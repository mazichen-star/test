"""Excel / Word 报告导出。"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd


def excel_bytes(
    data_out: pd.DataFrame,
    *,
    metrics: dict[str, Any] | None = None,
    model_name: str = "",
) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        data_out.to_excel(writer, index=False, sheet_name="预测结果")
        if metrics:
            info = pd.DataFrame(
                {
                    "项目": ["模型", "R²", "MAE", "RMSE", "5折交叉验证 R²"],
                    "数值": [
                        model_name,
                        metrics.get("r2", ""),
                        metrics.get("mae", ""),
                        metrics.get("rmse", ""),
                        metrics.get("cv_r2", ""),
                    ],
                }
            )
            info.to_excel(writer, index=False, sheet_name="训练信息")
    return buffer.getvalue()


def _set_doc_font(document, font_name: str = "Microsoft YaHei") -> None:
    try:
        from docx.oxml.ns import qn

        style = document.styles["Normal"]
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    except Exception:
        pass


def make_word_report(
    *,
    target_col: str,
    data_source: str,
    stats: dict[str, Any],
    feature_cols: list[str],
    model_name: str,
    metrics: dict[str, Any],
    cmp_df: pd.DataFrame | None,
    y_test: np.ndarray,
    y_pred: np.ndarray,
    result_df: pd.DataFrame | None = None,
    importance: pd.Series | None = None,
    generated_at: str | None = None,
    font_name: str | None = None,
) -> bytes | None:
    """生成 Word 训练报告；python-docx 未安装时返回 ``None``。"""

    try:
        from docx import Document
        from docx.shared import Inches
    except Exception:
        return None

    generated_at = generated_at or datetime.now().strftime("%Y-%m-%d %H:%M")
    document = Document()
    _set_doc_font(document)
    document.add_heading("材料力学性能预测训练报告", 0)
    document.add_paragraph(f"目标性能：{target_col}")
    document.add_paragraph(f"生成时间：{generated_at}")

    document.add_heading("1. 数据概况", 1)
    document.add_paragraph(f"数据文件：{data_source}")
    document.add_paragraph(
        f"原始行数：{stats.get('total_rows', '')}；有效样本：{stats.get('valid_rows', '')}"
    )
    document.add_paragraph(
        f"目标列缺失剔除：{stats.get('target_dropped', 0)} 行；"
        f"特征缺失剔除：{stats.get('feature_dropped', 0)} 行"
    )
    document.add_paragraph(f"特征列：{', '.join(feature_cols)}")

    document.add_heading("2. 模型与指标", 1)
    document.add_paragraph(f"最优模型：{model_name}")
    document.add_paragraph(f"测试集 R² = {metrics.get('r2', float('nan')):.4f}")
    document.add_paragraph(f"MAE = {metrics.get('mae', float('nan')):.4f}")
    document.add_paragraph(f"RMSE = {metrics.get('rmse', float('nan')):.4f}")
    cv_r2 = metrics.get("cv_r2")
    if cv_r2 is None or (isinstance(cv_r2, float) and np.isnan(cv_r2)):
        document.add_paragraph("5 折交叉验证 R²：样本不足或部分折无法计算")
    else:
        document.add_paragraph(f"5 折交叉验证 R² = {cv_r2:.4f}")

    if cmp_df is not None and len(cmp_df):
        document.add_paragraph("模型对比排名：")
        table = document.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "模型"
        table.rows[0].cells[1].text = "R²"
        table.rows[0].cells[2].text = "MAE"
        table.rows[0].cells[3].text = "RMSE"
        for _, row in cmp_df.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(row.get("模型", ""))
            cells[1].text = f"{row.get('R2', float('nan')):.4f}"
            cells[2].text = f"{row.get('MAE', float('nan')):.4f}"
            cells[3].text = f"{row.get('RMSE', float('nan')):.4f}"

    document.add_heading("3. 预测 vs 真实", 1)
    fig, ax = plt.subplots(figsize=(6, 5))
    try:
        if font_name:
            matplotlib.rcParams["font.sans-serif"] = [font_name, "DejaVu Sans"]
        matplotlib.rcParams["axes.unicode_minus"] = False
        ax.scatter(y_test, y_pred, alpha=0.7, color="#4C9AFF")
        low = float(min(np.min(y_test), np.min(y_pred)))
        high = float(max(np.max(y_test), np.max(y_pred)))
        ax.plot([low, high], [low, high], "r--", lw=1.5)
        ax.set_xlabel("真实值")
        ax.set_ylabel("预测值")
        ax.set_title(f"{model_name}  R²={metrics.get('r2', float('nan')):.3f}")
        image = io.BytesIO()
        fig.savefig(image, format="png", dpi=150, bbox_inches="tight")
        image.seek(0)
        document.add_picture(image, width=Inches(5))
    finally:
        plt.close(fig)

    if importance is not None and len(importance):
        document.add_heading("4. 特征重要性", 1)
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "特征"
        table.rows[0].cells[1].text = "重要性"
        for feature, value in importance.sort_values(ascending=False).items():
            cells = table.add_row().cells
            cells[0].text = str(feature)
            cells[1].text = f"{float(value):.4g}"

    if result_df is not None and len(result_df):
        document.add_heading("5. 预测结果明细（前 20 条）", 1)
        columns = ["真实值", "预测值", "误差"]
        if "预测标准差" in result_df.columns:
            columns.append("预测标准差")
        table = document.add_table(rows=1, cols=len(columns))
        for index, column in enumerate(columns):
            table.rows[0].cells[index].text = column
        for _, row in result_df.head(20).iterrows():
            cells = table.add_row().cells
            for index, column in enumerate(columns):
                value = row.get(column)
                try:
                    cells[index].text = f"{float(value):.4f}"
                except (TypeError, ValueError):
                    cells[index].text = "" if pd.isna(value) else str(value)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
